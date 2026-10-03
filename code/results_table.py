"""results_table.py — lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx (đừng gõ tay hàng chục dòng, rất dễ sai).

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, đừng ghi đè)

Mẫu dùng giá trị hiển thị (có danh sách chọn): loss CE|MSE, optimizer SGD|SGD+momentum|Adam|AdamW,
hidden "256-128", clip_norm 'none' hoặc số, diverged Y|N. to_row tự đổi từ giá trị trong cfg.
"""
from __future__ import annotations

import json
import math
from pathlib import Path

import openpyxl

COLUMNS = ("exp_id", "group", "description", "loss", "optimizer", "lr", "weight_decay", "batch", "epochs", "hidden",
           "dropout", "clip_norm", "precision", "init", "seed", "step0_loss", "best_val_loss", "best_epoch",
           "final_train_loss", "final_val_loss", "val_acc", "val_macro_f1", "time_per_epoch_s", "peak_mem_MB",
           "diverged", "eval_acc", "eval_macro_f1", "figure_file", "notes")
GROUPS = ("baseline", "loss", "optimizer", "hparam", "dropout", "clipping", "amp", "init", "final", "other")
LOSS_NAMES = {"ce": "CE", "mse": "MSE"}
OPT_NAMES = {"sgd": "SGD", "sgd_momentum": "SGD+momentum", "adam": "Adam", "adamw": "AdamW"}
SUMMARY_KEYS = ("step0_loss", "best_val_loss", "best_epoch", "final_train_loss", "final_val_loss", "val_acc",
                "val_macro_f1", "time_per_epoch_s", "peak_mem_MB")


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có."""
    out = Path(results_dir)
    out.mkdir(parents=True, exist_ok=True)
    cfg = {**result["cfg"], "hidden": list(result["cfg"]["hidden"])}
    path = out / f"{cfg['exp_id']}.json"
    with open(path, "w", encoding="utf-8") as f:
        json.dump({"cfg": cfg, "history": result["history"], "summary": result["summary"]}, f,
                  ensure_ascii=False, indent=1)
    return str(path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo nhóm như trong mẫu, rồi theo exp_id)."""
    results = []
    for p in Path(results_dir).glob("*.json"):
        with open(p, encoding="utf-8") as f:
            r = json.load(f)
        r["cfg"]["hidden"] = tuple(r["cfg"]["hidden"])
        results.append(r)
    order = {g: i for i, g in enumerate(GROUPS)}
    return sorted(results, key=lambda r: (order.get(r["cfg"]["group"], len(GROUPS)), r["cfg"]["exp_id"]))


def _num(v):
    """Số hữu hạn -> giữ nguyên; NaN/inf/None -> None (ô trống), vì xlsx không lưu được NaN."""
    if v is None or (isinstance(v, float) and not math.isfinite(v)):
        return None
    return v


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá trùng tên cột ở đầu file, giá trị theo danh sách chọn của mẫu.
    Chỉ truyền eval_scores (nội dung eval_result.json) cho baseline và cấu hình cuối cùng.
    notes: nếu rỗng thì lấy cfg["notes"]; tự thêm ghi chú về scheduler, phần cứng CPU và diverged.
    """
    cfg, s = result["cfg"], result["summary"]
    assert cfg["group"] in GROUPS, f"group phải là một trong {GROUPS}"
    extra = []
    if cfg.get("scheduler"):
        extra.append(f"scheduler={cfg['scheduler']}" + (f", warmup {cfg['warmup_steps']} bước" if cfg.get("warmup_steps") else ""))
    if s.get("device") in ("cpu", "mps"):
        extra.append(f"chạy trên {s['device']} (không đo bộ nhớ GPU)")
    if s.get("diverged"):
        extra.append(f"diverged ở epoch {s.get('epochs_run')}: loss NaN/inf, dừng sớm")
    note = "; ".join([n for n in [notes or cfg.get("notes", "")] + extra if n])

    row = dict(
        exp_id=cfg["exp_id"], group=cfg["group"], description=cfg.get("description", ""),
        loss=LOSS_NAMES[cfg["loss"]], optimizer=OPT_NAMES[cfg["optimizer"]],
        lr=cfg["lr"], weight_decay=cfg.get("weight_decay", 0.0), batch=cfg["batch"], epochs=cfg["epochs"],
        hidden="-".join(str(h) for h in cfg["hidden"]), dropout=cfg["dropout"],
        clip_norm="none" if cfg.get("clip_norm") is None else cfg["clip_norm"],
        precision=cfg["precision"], init=cfg["init"], seed=cfg["seed"],
        **{k: _num(s.get(k)) for k in SUMMARY_KEYS},
        diverged="Y" if s.get("diverged") else "N",
        eval_acc=_num(eval_scores["accuracy"]) if eval_scores else None,
        eval_macro_f1=_num(eval_scores["macro_f1"]) if eval_scores else None,
        figure_file=f"figures/{cfg['exp_id']}.png",
        notes=note,
    )
    assert tuple(row) == COLUMNS
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str, seed_ids: list[str] | None = None,
               summary_notes: dict | None = None) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    - Giữ công thức: không dùng data_only=True; không ghi vào cột công thức (AD..AG: step0_gap_vs_lnC,
      gap_val_minus_train, delta_val_f1_vs_base, beyond_noise). Dòng baseline mẫu (dòng 2) bị ghi đè/xoá.
    - seed_ids: exp_id các lần chạy baseline khác seed cho sheet "Seeds" (tối đa 5; mặc định giữ base-s1..3 của mẫu).
    - summary_notes: {group: nhận xét} ghi vào cột "nhận xét ngắn" của sheet "Summary".
    openpyxl không tính công thức: mở file bằng Excel/LibreOffice/Google Sheets và lưu lại để có giá trị.
    """
    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]
    header = {c.value: c.column for c in ws[1] if c.value is not None}
    formula_cols = {c.column for c in ws[2] if isinstance(c.value, str) and c.value.startswith("=")}
    input_cols = [header[k] for k in COLUMNS]
    assert not formula_cols & set(input_cols), "cột nhập trùng cột công thức"
    n_slots = sum(1 for r in range(2, ws.max_row + 1) if ws.cell(r, min(formula_cols)).value)
    if len(rows) > n_slots:
        raise ValueError(f"mẫu chỉ có công thức cho {n_slots} dòng, nhưng có {len(rows)} thí nghiệm")
    ids = [r["exp_id"] for r in rows]
    assert len(set(ids)) == len(ids), "exp_id phải duy nhất"

    for r in range(2, ws.max_row + 1):           # xoá dòng mẫu, giữ định dạng và công thức
        for col in input_cols:
            ws.cell(r, col).value = None
    for i, row in enumerate(rows, start=2):
        for k in COLUMNS:
            ws.cell(i, header[k]).value = row.get(k)

    if seed_ids is not None:
        assert len(seed_ids) <= 5, "sheet Seeds có 5 ô (A2:A6)"
        missing = set(seed_ids) - set(ids)
        assert not missing, f"seed_ids không có trong bảng: {missing}"
        seeds = wb["Seeds"]
        for j in range(5):
            seeds.cell(2 + j, 1).value = seed_ids[j] if j < len(seed_ids) else None

    if summary_notes:
        summ = wb["Summary"]
        for r in range(2, summ.max_row + 1):
            g = summ.cell(r, 1).value
            if g in summary_notes:
                summ.cell(r, 8).value = summary_notes[g]

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
